"""
Tests des livrables et des ajouts de la version 0.8.

Ce qui est vérifié ici, ce n'est pas que le code s'exécute — c'est que les
propriétés promises à l'utilisateur tiennent : un classeur qui contient
réellement des graphiques, une récupération EUR-Lex qui ne lève jamais, un
ciblage de négociation qui ne classe pas un opposant parmi les alliés.

Lancement :  python -m pytest tests/ -q
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.coalition import (best_additions, coalition_report,  # noqa: E402
                            negotiation_targets)
from core.eurlex import (en_pages, html_vers_texte, recuperer,  # noqa: E402
                         to_celex)
from core.excel import build_xlsx  # noqa: E402
from core.ingest import split_legal_text  # noqa: E402
from core.report import build_docx, build_fiche_docx  # noqa: E402
from core.scoring import ally_ranking, contentious_articles  # noqa: E402


# ---------------------------------------------------------------------------
# Jeu d'essai
# ---------------------------------------------------------------------------

def _matrice() -> pd.DataFrame:
    rng = np.random.default_rng(11)
    arts = [f"art_{i}" for i in range(1, 7)]
    ms = ["DE", "NL", "IT", "ES", "PL", "SE"]
    return pd.DataFrame(
        rng.choice([0, 1, 2, np.nan], size=(len(arts), len(ms)),
                   p=[0.2, 0.25, 0.35, 0.2]),
        index=arts, columns=ms)


def _detail() -> pd.DataFrame:
    return pd.DataFrame({
        "ms_code": ["DE", "NL", "IT"],
        "section_label": ["Article 1", "Article 2", "Article 1"],
        "section_order": [1, 2, 1],
        "stance": ["oppose", "aligne", "partiel"],
        "summary_fr": ["Demande la suppression.", "Soutient le texte.",
                       "Réserve d'examen."],
        "evidence": ["requests the deletion of this Article",
                     "supports the text as drafted",
                     "maintains a scrutiny reservation"],
        "scrutiny": [0, 0, 1], "deletion": [1, 0, 0],
        "themes": ['["gouvernance"]', "", None],
        "method": ["llm", "llm", "heuristique"],
        "confidence": [0.8, 0.9, 0.5],
        "page_start": [3, 7, 3],
        "text": ["Texte intégral 1.", "Texte intégral 2.", "Texte intégral 3."],
    })


# ---------------------------------------------------------------------------
# Classeur Excel
# ---------------------------------------------------------------------------

def test_classeur_contient_des_graphiques_natifs():
    """La demande était explicite : des graphiques, pas des tableaux de chiffres."""
    import io

    import openpyxl

    m = _matrice()
    octets = build_xlsx("Titre", "Dossier", m, ally_ranking(m),
                        contentious_articles(m), _detail())
    wb = openpyxl.load_workbook(io.BytesIO(octets))

    assert {"Lecture", "Matrice", "Classement EM", "Articles clivants",
            "Couverture", "Détail"} <= set(wb.sheetnames)
    graphiques = sum(len(wb[n]._charts) for n in wb.sheetnames)
    assert graphiques >= 3, "chaque feuille de données porte son graphique"


def test_classeur_porte_le_texte_de_detail():
    import io

    import openpyxl

    m = _matrice()
    octets = build_xlsx("Titre", "Dossier", m, ally_ranking(m),
                        contentious_articles(m), _detail())
    ws = openpyxl.load_workbook(io.BytesIO(octets))["Détail"]
    contenu = "\n".join(str(c.value) for row in ws.iter_rows() for c in row
                        if c.value)
    assert "requests the deletion of this Article" in contenu
    assert "Demande la suppression." in contenu


def test_classeur_supporte_un_detail_vide():
    """L'export doit rester possible avant toute analyse."""
    m = _matrice()
    octets = build_xlsx("Titre", "Dossier", m, ally_ranking(m),
                        contentious_articles(m), pd.DataFrame())
    assert len(octets) > 5000


# ---------------------------------------------------------------------------
# Documents Word
# ---------------------------------------------------------------------------

def test_note_word_contient_le_detail_et_ses_citations():
    import io

    from docx import Document

    m = _matrice()
    octets = build_docx("Titre", "Dossier", m, ally_ranking(m),
                        contentious_articles(m), detail=_detail())
    doc = Document(io.BytesIO(octets))
    texte = "\n".join(p.text for p in doc.paragraphs)
    assert "Détail des positions" in texte
    assert "requests the deletion of this Article" in texte
    assert "p. 3" in texte


def test_fiche_orientation_sans_element_reste_valide():
    import io

    from docx import Document

    octets = build_fiche_docx("Fiche", [("Acteur", "PME")], "", [], [],
                              ["Point non tranché."], ["RGPD"])
    doc = Document(io.BytesIO(octets))
    texte = "\n".join(p.text for p in doc.paragraphs)
    assert "Point non tranché." in texte


# ---------------------------------------------------------------------------
# EUR-Lex
# ---------------------------------------------------------------------------

def test_reference_traduite_en_celex():
    assert to_celex("32023R2854") == "32023R2854"
    assert to_celex("règlement (UE) 2023/2854") == "32023R2854"
    assert to_celex("directive (UE) 2019/1024") == "32019L1024"
    assert to_celex(
        "https://eur-lex.europa.eu/legal-content/FR/TXT/?uri=CELEX:32016R0679"
    ) == "32016R0679"
    assert to_celex("bonjour") == ""


def test_html_eurlex_decoupe_en_articles():
    """Le HTML doit produire les mêmes sections qu'un PDF, sans césure."""
    html = (
        "<html><head><title>Règlement</title><style>p{}</style></head><body>"
        "<p>(1) Considérant d'ouverture, suffisamment long pour compter.</p>"
        "<p class='oj-ti-art'>Article premier</p><p>Objet du règlement.</p>"
        "<p class='oj-ti-art'>Article 2</p><p>Définitions applicables.</p>"
        "<p class='oj-ti-art'>Article 5 bis</p><p>Disposition additionnelle.</p>"
        "<script>var x=1;</script></body></html>"
    )
    texte = html_vers_texte(html)
    assert "var x" not in texte, "le script ne doit pas entrer dans le texte"
    segments = split_legal_text(en_pages(texte, 200))
    ids = [s.section_id for s in segments]
    assert "art_1" in ids and "art_2" in ids and "art_5_bis" in ids


def test_recuperation_ne_leve_jamais():
    """Un poste filtré doit lire un message, pas une trace d'exception."""
    res = recuperer("référence incompréhensible")
    assert not res.ok and res.erreur

    res = recuperer("32023R2854", timeout=1)
    assert isinstance(res.erreur, str)   # succès ou échec, jamais d'exception


# ---------------------------------------------------------------------------
# Ciblage de négociation
# ---------------------------------------------------------------------------

def test_ciblage_ne_classe_pas_un_opposant_parmi_les_allies():
    m = pd.DataFrame(
        {"DE": [2, 2, 2], "NL": [0, 0, 1], "IT": [1, 1, 2]},
        index=["art_1", "art_2", "art_3"])
    cibles = negotiation_targets(m).set_index("EM")
    assert cibles.loc["DE", "Statut"] == "Allié"
    assert cibles.loc["NL", "Statut"] == "Opposé"
    assert cibles.loc["IT", "Statut"] in ("À convaincre", "Allié")


def test_ciblage_respecte_le_perimetre_d_articles():
    m = pd.DataFrame({"DE": [2, 0], "NL": [0, 2]}, index=["art_1", "art_2"])
    cibles = negotiation_targets(m, ["art_1"]).set_index("EM")
    assert cibles.loc["DE", "Articles alignés"] == 1
    assert cibles.loc["DE", "Divergences"] == 0


def test_minorite_de_blocage_exige_quatre_etats():
    """Trois grands États pèsent lourd et ne bloquent pas : c'est la règle."""
    trois = coalition_report(["FR", "DE", "IT"])
    if not trois["population_disponible"]:
        return
    assert trois["part_population"] > 0.35
    assert not trois["bloque"], "moins de quatre États ne bloquent jamais"
    quatre = coalition_report(["FR", "DE", "IT", "ES"])
    assert quatre["bloque"]


def test_ralliements_classes_par_apport_de_population():
    gains = best_additions(["FR", "DE"], ["MT", "ES", "LU"], "blocage")
    if gains.empty:
        return
    assert list(gains["EM"])[0] == "ES"


# ---------------------------------------------------------------------------
# Le classeur reprend la forme du suivi tenu à la main
# ---------------------------------------------------------------------------

def _classeur_suivi():
    import io

    import openpyxl

    from core import palette

    palette.utiliser("note")
    m = _matrice()
    octets = build_xlsx(
        "Titre", "Dossier", m, ally_ranking(m), contentious_articles(m),
        _detail(),
        # Les clés sont les intitulés de la matrice — ce que fait la page.
        fr_positions={"art_1": "Étendre le champ aux technologies sensibles."},
        sujets={"art_1": "Champ du cadre"})
    return openpyxl.load_workbook(io.BytesIO(octets))


def test_ordre_des_feuilles_suit_le_suivi_a_la_main():
    """Position française, puis enjeux, puis détail : l'ordre de lecture."""
    noms = _classeur_suivi().sheetnames
    assert noms[:4] == ["Lecture", "Position FR", "Synthèse thématique",
                        "Détail par article"]
    # L'onglet thématique passe avant le détail article par article.
    assert noms.index("Synthèse thématique") < noms.index("Détail par article")


def test_detail_par_article_porte_la_position_francaise_en_colonne_b():
    ws = _classeur_suivi()["Détail par article"]
    entete = [ws.cell(row=4, column=j).value for j in range(1, 5)]
    assert entete[0] == "Article"
    assert entete[1] and entete[1].startswith("FR")

    colonne_a = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)]
    assert any("Champ du cadre" in str(v) for v in colonne_a), \
        "le sujet de l'article doit figurer sous son numéro"
    colonne_b = [ws.cell(row=r, column=2).value for r in range(5, ws.max_row + 1)]
    assert any("Étendre le champ" in str(v) for v in colonne_b)
    assert any("règle du silence" in str(v) for v in colonne_b), \
        "les articles sans amendement français doivent le dire"


def test_cases_du_suivi_sont_colorees():
    """Le code couleur est ce qui rend le tableau lisible d'un coup d'œil."""
    from core.palette import PALETTES

    ws = _classeur_suivi()["Détail par article"]
    attendues = {PALETTES["note"][k].lstrip("#").upper()
                 for k in ("aligne", "partiel", "oppose", "absent")}
    trouvees = set()
    for row in ws.iter_rows(min_row=5, min_col=3):
        for c in row:
            if c.fill and c.fill.fill_type == "solid":
                trouvees.add(str(c.fill.fgColor.rgb)[-6:].upper())
    assert trouvees & attendues, "aucune case n'est colorée"


def test_synthese_thematique_regroupe_par_enjeu():
    ws = _classeur_suivi()["Synthèse thématique"]
    contenu = "\n".join(str(c.value) for row in ws.iter_rows() for c in row
                        if c.value)
    assert "Gouvernance" in contenu or "gouvernance" in contenu


def test_colonne_themes_avant_les_colonnes_techniques():
    """Demande explicite : le thème doit être visible sans faire défiler."""
    from core.excel import COLONNES_DETAIL

    ordre = [c for c, _, _ in COLONNES_DETAIL]
    assert ordre.index("themes_fr") < ordre.index("evidence")
    assert ordre.index("themes_fr") < ordre.index("method")

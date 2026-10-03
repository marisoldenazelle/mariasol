"""
Livrables : images PNG et note Word modifiable.

Deux choix techniques, tous deux dictés par le poste de destination.

`matplotlib` plutôt que l'export d'image de Plotly : Plotly passe par un
navigateur Chrome pour produire un PNG, dépendance lourde et souvent absente
d'un poste d'administration. matplotlib écrit le fichier directement.

`python-docx` plutôt qu'un moteur externe : le document doit se générer sur la
machine de l'utilisateur, hors ligne, et s'ouvrir dans Word en restant
entièrement modifiable — tableaux éditables, images repositionnables.
"""

from __future__ import annotations

import io
from datetime import date

import matplotlib
import numpy as np
import pandas as pd

matplotlib.use("Agg")            # aucun affichage : on écrit des fichiers
import matplotlib.patches as mpatches          # noqa: E402
import matplotlib.pyplot as plt                # noqa: E402
from matplotlib.colors import ListedColormap   # noqa: E402

INK = "#0b0b0b"
INK_SECONDARY = "#52514e"
INK_MUTED = "#898781"
GRID = "#e1e0d9"
SURFACE = "#ffffff"

# Les deux conventions de couleur vivent dans `core/palette.py` : elles
# doivent être identiques à l'écran, dans le PNG, dans le Word et dans
# l'Excel, sans quoi la même matrice change de sens selon le support.
from .palette import PALETTES, SYMBOLES, actif  # noqa: E402

plt.rcParams.update({
    "font.family": "sans-serif",
    "font.sans-serif": ["DejaVu Sans", "Segoe UI", "Arial"],
    "axes.edgecolor": GRID,
    "figure.facecolor": SURFACE,
    "axes.facecolor": SURFACE,
})


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight",
                facecolor=SURFACE)
    plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------

def heatmap_png(matrix: pd.DataFrame, titre: str = "",
                palette: str | None = None, labels: dict | None = None) -> bytes:
    """Matrice articles × États membres, avec symbole ET couleur dans chaque case."""
    if matrix.empty:
        raise ValueError("matrice vide")
    pal = PALETTES[palette or actif()]
    labels = labels or {}
    cmap = ListedColormap([pal["oppose"], pal["partiel"], pal["aligne"]])

    lignes = [labels.get(i, i) for i in matrix.index]
    h = max(3.0, 0.42 * len(matrix.index) + 1.6)
    w = max(6.0, 0.62 * len(matrix.columns) + 4.5)
    fig, ax = plt.subplots(figsize=(w, h))

    data = matrix.values.astype(float)
    ax.imshow(np.ma.masked_invalid(data), cmap=cmap, vmin=0, vmax=2,
              aspect="auto")
    # Les cases sans contribution portent leur propre teinte, pas un blanc
    # qui se confondrait avec le fond de la page.
    for y in range(data.shape[0]):
        for x in range(data.shape[1]):
            if np.isnan(data[y, x]):
                ax.add_patch(mpatches.Rectangle(
                    (x - 0.5, y - 0.5), 1, 1, facecolor=pal["absent"],
                    edgecolor=SURFACE, linewidth=1.5))
            else:
                v = int(data[y, x])
                couleur = INK if (v == 1 or not pal["texte_fonce"]) else "#ffffff"
                ax.text(x, y, SYMBOLES[v], ha="center", va="center",
                        fontsize=9, color=couleur)

    ax.set_xticks(range(len(matrix.columns)))
    ax.set_xticklabels(matrix.columns, fontsize=9, color=INK_SECONDARY)
    ax.set_yticks(range(len(lignes)))
    ax.set_yticklabels(lignes, fontsize=9, color=INK_SECONDARY)
    ax.set_xticks(np.arange(-0.5, len(matrix.columns), 1), minor=True)
    ax.set_yticks(np.arange(-0.5, len(lignes), 1), minor=True)
    ax.grid(which="minor", color=SURFACE, linewidth=1.5)
    ax.tick_params(which="minor", length=0)
    for spine in ax.spines.values():
        spine.set_visible(False)

    ax.legend(
        handles=[
            mpatches.Patch(color=pal["aligne"], label="Aligné (+)"),
            mpatches.Patch(color=pal["partiel"], label="Partiel (~)"),
            mpatches.Patch(color=pal["oppose"], label="Divergent (−)"),
            mpatches.Patch(color=pal["absent"], label="Pas de contribution"),
        ],
        loc="upper left", bbox_to_anchor=(1.01, 1.0), frameon=True,
        fontsize=9, edgecolor=GRID)

    if titre:
        ax.set_title(titre, fontsize=12, color=INK, pad=14, loc="left")
    return _png(fig)


def ranking_png(ranking: pd.DataFrame, titre: str = "",
                surligner: str = "FR") -> bytes:
    """Classement des États membres, l'État de référence mis en évidence."""
    if ranking.empty:
        raise ValueError("classement vide")
    d = ranking.sort_values("Score moyen", ascending=False)
    couleurs = ["#17427d" if i == surligner else "#8fb4e3" for i in d.index]

    fig, ax = plt.subplots(figsize=(max(6.0, 0.45 * len(d) + 2), 3.6))
    barres = ax.bar(range(len(d)), d["Score moyen"], color=couleurs,
                    edgecolor=SURFACE, linewidth=1.2, width=0.72)
    ax.bar_label(barres, fmt="%.2f", fontsize=8, color=INK_SECONDARY, padding=2)
    ax.set_xticks(range(len(d)))
    ax.set_xticklabels(d.index, fontsize=9, color=INK_SECONDARY)
    ax.set_ylabel("Score moyen (0 divergent → 2 aligné)", fontsize=9,
                  color=INK_MUTED)
    ax.set_ylim(0, 2.25)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    if titre:
        ax.set_title(titre, fontsize=12, color=INK, pad=12, loc="left")
    return _png(fig)


def engagement_png(df: pd.DataFrame, titre: str = "",
                   surligner: str = "FR") -> bytes:
    """Nombre de contributions par État membre — lecture de la mobilisation."""
    counts = df.groupby("ms_code").size().sort_values(ascending=False)
    if counts.empty:
        raise ValueError("aucune contribution")
    couleurs = ["#17427d" if i == surligner else "#8fb4e3" for i in counts.index]

    fig, ax = plt.subplots(figsize=(max(6.0, 0.45 * len(counts) + 2), 3.4))
    barres = ax.bar(range(len(counts)), counts.values, color=couleurs,
                    edgecolor=SURFACE, linewidth=1.2, width=0.72)
    ax.bar_label(barres, fontsize=8, color=INK_SECONDARY, padding=2)
    ax.set_xticks(range(len(counts)))
    ax.set_xticklabels(counts.index, fontsize=9, color=INK_SECONDARY)
    ax.set_ylabel("Contributions écrites", fontsize=9, color=INK_MUTED)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    if titre:
        ax.set_title(titre, fontsize=12, color=INK, pad=12, loc="left")
    return _png(fig)


# ---------------------------------------------------------------------------
# Note Word
# ---------------------------------------------------------------------------

def _fixed_layout(table) -> None:
    """Force une mise en page à largeurs fixes.

    Sans cela Word et LibreOffice recalculent les largeurs au rendu et
    l'intitulé d'article se coupe en trois lignes, quelles que soient les
    largeurs posées sur les cellules.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    el = OxmlElement("w:tblLayout")
    el.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(el)


def _shade(cell, hexa: str) -> None:
    """Applique un fond de cellule — python-docx n'expose pas l'ombrage."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexa.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(el)


def build_fiche_docx(
    titre: str,
    situation: list[tuple[str, str]],
    synthese: str,
    elements: list[dict],
    contexte: list[dict] | None = None,
    manques: list[str] | None = None,
    corpus: list[str] | None = None,
) -> bytes:
    """Fiche d'orientation Word — une question, ce que les textes en disent.

    Chaque élément porte sa citation littérale et sa référence. C'est ce qui
    distingue la fiche d'une réponse de chatbot : le destinataire peut la
    contrôler ligne à ligne dans le texte officiel, et la reprendre à son
    compte en la modifiant.
    """
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for marge in ("left_margin", "right_margin"):
        setattr(doc.sections[0], marge, Cm(2.2))

    doc.add_heading(titre, level=0)
    p = doc.add_paragraph(f"Fiche établie le {date.today().strftime('%d/%m/%Y')}")
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    if situation:
        doc.add_heading("Situation examinée", level=1)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for cle, valeur in situation:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = cle, valeur
            cells[0].width, cells[1].width = Cm(5.2), Cm(11.0)
            for r in cells[0].paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(9)
            for r in cells[1].paragraphs[0].runs:
                r.font.size = Pt(9)
        _fixed_layout(t)

    if synthese:
        doc.add_heading("Ce que disent les textes chargés", level=1)
        doc.add_paragraph(synthese)

    def _bloc(items: list[dict], titre_bloc: str, note: str = "") -> None:
        if not items:
            return
        doc.add_heading(titre_bloc, level=1)
        if note:
            n = doc.add_paragraph(note)
            n.runs[0].font.size = Pt(9)
            n.runs[0].italic = True
            n.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
        for i, el in enumerate(items, start=1):
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(2)
            tete = par.add_run(f"{i}. ")
            tete.bold = True
            corps = par.add_run(str(el.get("statement", "")).strip())
            corps.font.size = Pt(10.5)

            q = doc.add_paragraph(f"« {str(el.get('quote', '')).strip()} »")
            q.paragraph_format.left_indent = Cm(0.8)
            q.paragraph_format.space_after = Pt(2)
            q.runs[0].italic = True
            q.runs[0].font.size = Pt(9)

            src = doc.add_paragraph(str(el.get("citation", "")).strip())
            src.paragraph_format.left_indent = Cm(0.8)
            src.paragraph_format.space_after = Pt(8)
            src.runs[0].font.size = Pt(8)
            src.runs[0].font.color.rgb = RGBColor(0x89, 0x87, 0x81)

    _bloc(elements, "Dispositions applicables")
    _bloc(contexte or [], "Éléments de contexte",
          "Ces passages n'apportent pas de réponse directe mais éclairent la "
          "situation examinée.")

    if manques:
        doc.add_heading("Points que les textes chargés ne tranchent pas", level=1)
        for m in manques:
            doc.add_paragraph(str(m), style="List Bullet")

    if corpus:
        doc.add_heading("Textes interrogés", level=1)
        for nom in corpus:
            c = doc.add_paragraph(str(nom), style="List Bullet")
            c.runs[0].font.size = Pt(9)

    doc.add_paragraph()
    avert = doc.add_paragraph(
        "Orientation documentaire de premier niveau, produite à partir des "
        "seuls textes chargés dans l'outil. Toute affirmation dont la citation "
        "n'a pas été retrouvée littéralement dans le texte source a été "
        "retirée avant production. Ce document n'est pas un conseil juridique "
        "et ne dispense pas d'une analyse au cas d'espèce.")
    avert.runs[0].font.size = Pt(8)
    avert.runs[0].italic = True
    avert.runs[0].font.color.rgb = RGBColor(0x89, 0x87, 0x81)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


STANCE_MOT = {"aligne": "Aligné", "partiel": "Partiel", "oppose": "Divergent",
              "neutre": "Neutre"}
STANCE_SIGNE = {"aligne": "+", "partiel": "~", "oppose": "−"}


def _section_detail(doc, detail: pd.DataFrame, labels: dict, pal: dict,
                    detail_max: int, fr_positions: dict | None = None,
                    sujets: dict | None = None) -> None:
    """Le texte de chaque contribution, groupé par article.

    Un paragraphe par État membre : position, résumé, puis la citation
    littérale retenue et sa page. C'est ce qui rend la note opposable — le
    lecteur peut vérifier chaque ligne dans le document source sans l'outil.
    """
    import json

    from docx.shared import Cm, Pt, RGBColor

    from .labels import reference_dans_texte

    doc.add_heading("Détail des positions, article par article", level=1)
    intro = doc.add_paragraph(
        "Pour chaque article : la position retenue, ce qu'en dit l'État membre, "
        "et le passage du document source qui le fonde.")
    intro.runs[0].font.size = Pt(9)
    intro.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    d = detail.copy()
    if "section_order" in d.columns:
        d = d.sort_values(["section_order", "ms_code"])
    else:
        d = d.sort_values(["section_label", "ms_code"])

    tronque = len(d) > detail_max
    d = d.head(detail_max)

    fr_positions = fr_positions or {}
    sujets = sujets or {}

    for label, groupe in d.groupby("section_label", sort=False):
        titre_art = labels.get(label, label)
        sujet = str(sujets.get(label, "") or "").strip()
        doc.add_heading(str(titre_art) + (f" · {sujet}" if sujet else ""), level=2)

        # La position française d'abord : « DE est opposé » n'a de sens que si
        # l'on sait opposé à quoi.
        reference = str(fr_positions.get(label, "") or "").strip()
        ref_par = doc.add_paragraph()
        ref_par.paragraph_format.space_after = Pt(6)
        etiquette = ref_par.add_run("Position française de référence, ")
        etiquette.bold = True
        etiquette.font.size = Pt(9.5)
        corps = ref_par.add_run(
            reference or "aucun amendement français : maintien du texte en "
                         "l'état (règle du silence).")
        corps.font.size = Pt(9.5)
        corps.italic = not reference

        for _, row in groupe.iterrows():
            stance = row.get("stance") or ""
            disposition = reference_dans_texte(str(row.get("text") or ""))
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            tete = p.add_run(f"{row['ms_code']} · "
                             + (f"{disposition} · " if disposition else "")
                             + f"{STANCE_MOT.get(stance, 'non classé')} "
                             f"{STANCE_SIGNE.get(stance, '')}  ")
            tete.bold = True
            tete.font.size = Pt(9.5)
            tete.font.color.rgb = RGBColor.from_string(
                {"aligne": pal["aligne"], "partiel": INK_SECONDARY,
                 "oppose": pal["oppose"]}.get(stance, INK_SECONDARY
                                              ).lstrip("#").upper())
            resume = p.add_run(str(row.get("summary_fr") or "").strip())
            resume.font.size = Pt(9.5)

            preuve = str(row.get("evidence") or "").strip()
            if preuve:
                q = doc.add_paragraph(f"« {preuve} »")
                q.paragraph_format.left_indent = Cm(0.8)
                q.paragraph_format.space_after = Pt(2)
                q.runs[0].italic = True
                q.runs[0].font.size = Pt(8.5)
                q.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

            meta = []
            if row.get("scrutiny"):
                meta.append("réserve d'examen")
            if row.get("deletion"):
                meta.append("demande de suppression")
            try:
                themes = json.loads(row.get("themes") or "[]")
            except (TypeError, ValueError):
                themes = []
            if themes:
                meta.append("thèmes : " + ", ".join(themes))
            page = row.get("page_start")
            if pd.notna(page):
                meta.append(f"p. {int(page)}")
            if row.get("method"):
                meta.append(str(row["method"]))
            if meta:
                m = doc.add_paragraph(" · ".join(meta))
                m.paragraph_format.left_indent = Cm(0.8)
                m.paragraph_format.space_after = Pt(6)
                m.runs[0].font.size = Pt(8)
                m.runs[0].font.color.rgb = RGBColor(0x89, 0x87, 0x81)

    if tronque:
        avis = doc.add_paragraph(
            f"Détail limité aux {detail_max} premières contributions. "
            "Restreignez le périmètre de l'export pour une note complète, ou "
            "utilisez le classeur Excel, qui les porte toutes.")
        avis.runs[0].font.size = Pt(8)
        avis.runs[0].italic = True


def build_docx(
    titre: str,
    dossier: str,
    matrix: pd.DataFrame,
    ranking: pd.DataFrame,
    contentious: pd.DataFrame,
    images: list[tuple[str, bytes]] | None = None,
    labels: dict | None = None,
    palette: str | None = None,
    methodologie: str = "",
    detail: pd.DataFrame | None = None,
    detail_max: int = 400,
    fr_positions: dict | None = None,
    sujets: dict | None = None,
) -> bytes:
    """Produit une note Word entièrement modifiable.

    Les tableaux sont de vrais tableaux Word — pas des images — pour que la
    note puisse être reprise, annotée et complétée sans repasser par l'outil.
    """
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    pal = PALETTES[palette or actif()]
    labels = labels or {}
    doc = Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for marge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, marge, Cm(1.6))

    doc.add_heading(titre, level=0)
    p = doc.add_paragraph()
    run = p.add_run(
        f"{dossier} · établie le {date.today().strftime('%d/%m/%Y')} · "
        f"{len(matrix.columns)} États membres · {len(matrix.index)} articles")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    if methodologie:
        doc.add_heading("Méthode", level=1)
        m = doc.add_paragraph(methodologie)
        m.runs[0].font.size = Pt(9.5)

    # --- matrice ----------------------------------------------------------
    doc.add_heading("Positions article par article", level=1)
    table = doc.add_table(rows=1, cols=len(matrix.columns) + 1)
    table.style = "Table Grid"
    entete = table.rows[0].cells
    entete[0].text = "Article"
    for j, col in enumerate(matrix.columns, start=1):
        entete[j].text = str(col)
    for cell in entete:
        for par in cell.paragraphs:
            for r in par.runs:
                r.bold = True
                r.font.size = Pt(8)
        _shade(cell, "#f0efec")

    # Largeurs explicites : sans cela Word répartit à parts égales et
    # « Commentaires généraux » se coupe en trois lignes dans la colonne des
    # intitulés. La largeur doit être posée sur chaque cellule, pas seulement
    # sur la colonne.
    largeur_intitule = Cm(4.6)
    largeur_etat = Cm(max(0.9, min(1.4, 21.0 / max(len(matrix.columns), 1))))
    entete[0].width = largeur_intitule
    for j in range(1, len(entete)):
        entete[j].width = largeur_etat
    table.autofit = False
    _fixed_layout(table)
    # La largeur doit être posée sur la colonne ET sur chaque cellule : Word
    # lit la grille, LibreOffice lit les cellules.
    table.columns[0].width = largeur_intitule
    for j in range(1, len(table.columns)):
        table.columns[j].width = largeur_etat

    for idx, row in matrix.iterrows():
        cells = table.add_row().cells
        cells[0].text = labels.get(idx, str(idx))
        for r in cells[0].paragraphs[0].runs:
            r.font.size = Pt(8)
        cells[0].width = largeur_intitule
        for j in range(1, len(cells)):
            cells[j].width = largeur_etat
        for j, col in enumerate(matrix.columns, start=1):
            v = row[col]
            if pd.isna(v):
                _shade(cells[j], pal["absent"])
                continue
            v = int(v)
            cells[j].text = SYMBOLES[v]
            par = cells[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in par.runs:
                r.font.size = Pt(9)
                r.bold = True
            _shade(cells[j], {2: pal["aligne"], 1: pal["partiel"],
                              0: pal["oppose"]}[v])

    leg = doc.add_paragraph("+ aligné  ·  ~ partiel  ·  − divergent  ·  "
                            "case vide : pas de contribution")
    leg.runs[0].font.size = Pt(8)
    leg.runs[0].font.color.rgb = RGBColor(0x89, 0x87, 0x81)

    # --- classement --------------------------------------------------------
    if not ranking.empty:
        doc.add_heading("Proximité avec la position française", level=1)
        t = doc.add_table(rows=1, cols=6)
        t.style = "Table Grid"
        for j, h in enumerate(["EM", "Score moyen", "Articles analysés",
                               "Alignements", "Partiels", "Divergences"]):
            t.rows[0].cells[j].text = h
            for r in t.rows[0].cells[j].paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(9)
            _shade(t.rows[0].cells[j], "#f0efec")
        for code, row in ranking.iterrows():
            cells = t.add_row().cells
            valeurs = [code, f"{row['Score moyen']:.2f}",
                       row["Articles analysés"], row["Alignements"],
                       row["Partiels"], row["Oppositions"]]
            for j, v in enumerate(valeurs):
                cells[j].text = str(v)
                for r in cells[j].paragraphs[0].runs:
                    r.font.size = Pt(9)

    # --- articles clivants -------------------------------------------------
    if not contentious.empty:
        doc.add_heading("Articles les plus clivants", level=1)
        top = contentious.head(10)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for j, h in enumerate(["Article", "Score moyen", "Dispersion",
                               "Divergences"]):
            t.rows[0].cells[j].text = h
            for r in t.rows[0].cells[j].paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(9)
            _shade(t.rows[0].cells[j], "#f0efec")
        for idx, row in top.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate([labels.get(idx, str(idx)),
                                   f"{row['Score moyen']:.2f}",
                                   f"{row['Dispersion']:.2f}"
                                   if pd.notna(row["Dispersion"]) else "—",
                                   row["Oppositions"]]):
                cells[j].text = str(v)
                for r in cells[j].paragraphs[0].runs:
                    r.font.size = Pt(9)

    # --- figures -----------------------------------------------------------
    for legende, png in (images or []):
        doc.add_page_break()
        doc.add_heading(legende, level=1)
        doc.add_picture(io.BytesIO(png), width=Cm(24))

    # --- détail article par article ----------------------------------------
    # La matrice dit « DE diverge sur l'article 7 ». Elle ne dit pas pourquoi,
    # et c'est le pourquoi qu'un chef de bureau demande. Cette partie porte le
    # texte : ce que dit l'État membre, la citation qui le prouve, la page.
    #
    # Elle bascule en portrait : une matrice se lit en largeur, un paragraphe
    # se lit sur une colonne étroite. Word gère les deux dans un même fichier.
    if detail is not None and not detail.empty:
        from docx.enum.section import WD_SECTION


        portrait = doc.add_section(WD_SECTION.NEW_PAGE)
        portrait.orientation = WD_ORIENT.PORTRAIT
        portrait.page_width, portrait.page_height = (
            min(portrait.page_width, portrait.page_height),
            max(portrait.page_width, portrait.page_height))
        for marge in ("left_margin", "right_margin"):
            setattr(portrait, marge, Cm(2.2))
        _section_detail(doc, detail, labels, pal, detail_max,
                        fr_positions, sujets)

    doc.add_paragraph()
    avert = doc.add_paragraph(
        "Document produit automatiquement à partir des contributions écrites "
        "des États membres. Le classement de chaque position est adossé à une "
        "citation vérifiée dans le document source ; il reste à valider par "
        "l'analyste avant tout usage en négociation.")
    avert.runs[0].font.size = Pt(8)
    avert.runs[0].italic = True
    avert.runs[0].font.color.rgb = RGBColor(0x89, 0x87, 0x81)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Note de synthèse sur un acte modificatif
# ---------------------------------------------------------------------------

_COULEUR_IMPACT = {"majeur": "#f4cccc", "mineur": "#fff2cc",
                   "redactionnel": "#f0efec", "": "#ffffff"}


def note_omnibus_docx(nom_omnibus: str, synthese, version: str = "") -> bytes:
    """Note de synthèse Word sur un omnibus, prête à être diffusée.

    Ce que l'agent produit à la fin, ce n'est pas un écran, c'est une note.
    Elle est donc entièrement modifiable : de vrais tableaux Word, pas des
    images, pour qu'elle se reprenne et s'annote sans repasser par l'outil.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for marge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(doc.sections[0], marge, Cm(2.0))

    doc.add_heading(f"Ce que change {nom_omnibus}", level=0)
    p = doc.add_paragraph()
    run = p.add_run(
        f"Note de synthèse établie le {date.today().strftime('%d/%m/%Y')} · "
        f"{synthese.nb_textes} textes modifiés · {synthese.nb_articles} "
        f"articles touchés · {synthese.nb_instructions} instructions de "
        "modification")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    # --- en bref -----------------------------------------------------------
    doc.add_heading("En bref", level=1)
    majeurs = synthese.points_durs
    if synthese.caracterises:
        phrase = (
            f"L'acte modifie {synthese.nb_textes} textes et touche "
            f"{synthese.nb_articles} articles. Sur les "
            f"{synthese.caracterises} articles caractérisés, "
            f"{len(majeurs)} portent une modification de fond : obligation, "
            "champ d'application, seuil, compétence, sanction ou délai. Ce "
            "sont eux qui figurent en tête de la section suivante.")
    else:
        phrase = (
            f"L'acte modifie {synthese.nb_textes} textes et touche "
            f"{synthese.nb_articles} articles. La portée des modifications "
            "n'a pas été caractérisée : la note liste ce que l'omnibus fait, "
            "sans dire ce que cela pèse.")
    doc.add_paragraph(phrase)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, titre in enumerate(("Texte modifié", "Article de l'acte",
                               "Articles touchés", "Modifications")):
        cell = table.rows[0].cells[i]
        cell.text = titre
        for par in cell.paragraphs:
            for r in par.runs:
                r.bold = True
                r.font.size = Pt(9)
        _shade(cell, "#f0efec")
    for acte, lignes in synthese.par_texte.items():
        cells = table.add_row().cells
        cells[0].text = lignes[0].libelle_cible
        cells[1].text = lignes[0].article_omnibus
        cells[2].text = str(len(lignes))
        cells[3].text = str(sum(l.nb_modifications for l in lignes))
        for cell in cells:
            for par in cell.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(9)

    # --- les points durs ---------------------------------------------------
    if majeurs:
        doc.add_heading("Les modifications de fond", level=1)
        doc.add_paragraph(
            "Portée qualifiée par le modèle à partir des instructions "
            "écrites dans l'acte. Chaque entrée porte l'instruction verbatim "
            "et sa page : la qualification se conteste sur pièce.")
        for ligne in majeurs:
            titre = doc.add_paragraph()
            r = titre.add_run(f"{ligne.libelle_cible} · {ligne.article}")
            r.bold = True
            r.font.size = Pt(10.5)
            if ligne.resume:
                doc.add_paragraph(ligne.resume)
            # L'instruction et le texte qu'elle introduit vont ensemble :
            # « les points suivants sont ajoutés: » ne dit rien seul.
            nouveaux = list(ligne.textes_nouveaux) + [""] * len(ligne.instructions)
            for instruction, nouveau in list(zip(ligne.instructions,
                                                 nouveaux))[:6]:
                puce = doc.add_paragraph(instruction[:900], style="List Bullet")
                for r in puce.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
                if nouveau:
                    cite = doc.add_paragraph(f"« {nouveau[:1500]} »")
                    cite.paragraph_format.left_indent = Cm(1.4)
                    for r in cite.runs:
                        r.font.size = Pt(8.5)
                        r.italic = True
                        r.font.color.rgb = RGBColor(0x3A, 0x3A, 0x3A)
            if ligne.concepts:
                c = doc.add_paragraph()
                run = c.add_run("Notions touchées : " + ", ".join(ligne.concepts))
                run.font.size = Pt(9)
                run.italic = True

    # --- notions transverses ----------------------------------------------
    if synthese.notions:
        doc.add_heading("Ce qui revient d'un texte à l'autre", level=1)
        doc.add_paragraph(
            "Notions modifiées dans plusieurs textes à la fois. C'est la "
            "lecture horizontale d'un omnibus : elle dit quels bureaux "
            "traitent du même sujet dans des dossiers différents.")
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        for i, titre in enumerate(("Notion", "Articles concernés")):
            cell = t.rows[0].cells[i]
            cell.text = titre
            for par in cell.paragraphs:
                for r in par.runs:
                    r.bold = True
                    r.font.size = Pt(9)
            _shade(cell, "#f0efec")
        for notion, lignes in synthese.notions[:20]:
            cells = t.add_row().cells
            cells[0].text = notion
            cells[1].text = " · ".join(
                f"{l.nom_cible or l.acte_cible} {l.article}" for l in lignes[:8])
            for cell in cells:
                for par in cell.paragraphs:
                    for r in par.runs:
                        r.font.size = Pt(9)

    # --- le détail ---------------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Le détail, article par article", level=1)
    detail = doc.add_table(rows=1, cols=5)
    detail.style = "Table Grid"
    for i, titre in enumerate(("Texte", "Article", "Opérations", "Portée",
                               "Ce que ça change")):
        cell = detail.rows[0].cells[i]
        cell.text = titre
        for par in cell.paragraphs:
            for r in par.runs:
                r.bold = True
                r.font.size = Pt(8.5)
        _shade(cell, "#f0efec")
    largeurs = (Cm(3.6), Cm(2.2), Cm(2.6), Cm(1.8), Cm(6.5))
    for ligne in synthese.lignes:
        cells = detail.add_row().cells
        valeurs = (ligne.nom_cible or ligne.acte_cible, ligne.article,
                   ", ".join(ligne.operations), ligne.impact or "—",
                   ligne.resume or "—")
        for cell, valeur, largeur in zip(cells, valeurs, largeurs):
            cell.text = str(valeur)
            cell.width = largeur
            for par in cell.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(8.5)
        _shade(cells[3], _COULEUR_IMPACT.get(ligne.impact, "#ffffff"))
    detail.autofit = False
    _fixed_layout(detail)
    for col, largeur in zip(detail.columns, largeurs):
        col.width = largeur

    pied = doc.add_paragraph()
    run = pied.add_run(
        "Lecture des instructions : déterministe, aucune interprétation. "
        "Qualification de la portée : produite par un modèle de langage à "
        "partir des instructions écrites, à vérifier avant tout usage en "
        "négociation." + (f" Produit par mariasol {version}." if version else ""))
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x89, 0x87, 0x81)
    pied.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

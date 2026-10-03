"""
Classeur Excel de restitution — chiffres, graphiques et texte de détail.

Trois exigences ont dicté la construction.

*Des graphiques, pas seulement des tableaux de chiffres.* Deux techniques
cohabitent ici, et le choix n'est pas neutre :

- les **graphiques natifs Excel** (`openpyxl.chart`) restent liés aux cellules.
  L'agent peut filtrer, corriger une valeur, ajouter un État membre : le
  graphique suit. C'est ce qu'on veut pour un document de travail qui va
  circuler et être repris ;
- les **images** produites par matplotlib rendent ce qu'Excel ne sait pas
  faire proprement — la matrice article × État membre en damier coloré.

*Un classeur qui se lit sans mode d'emploi.* La première feuille explique la
méthode et la légende ; les feuilles de données ont des filtres, des volets
figés et des largeurs de colonnes posées.

*Le texte, pas seulement les scores.* La feuille « Détail » porte le résumé,
la citation vérifiée et la page de chaque contribution : c'est ce qui permet
de contester une case de la matrice sans rouvrir le PDF d'origine.
"""

from __future__ import annotations

import io
from datetime import date

import pandas as pd

from .palette import PALETTES, SYMBOLES, actif

# ---------------------------------------------------------------------------
# Constantes de mise en forme
# ---------------------------------------------------------------------------

TITRE_FILL = "F0EFEC"
ENTETE_FONT = "0B0B0B"


def _hex(couleur: str) -> str:
    return couleur.lstrip("#").upper()


def _ecrire_df(ws, df: pd.DataFrame, index_titre: str | None = None,
               depart: int = 1) -> int:
    """Écrit un tableau à partir de la ligne `depart`. Renvoie la ligne suivante."""
    from openpyxl.styles import Alignment, Font, PatternFill

    entete = ([index_titre] if index_titre else []) + [str(c) for c in df.columns]
    for j, valeur in enumerate(entete, start=1):
        cell = ws.cell(row=depart, column=j, value=valeur)
        cell.font = Font(bold=True, color=ENTETE_FONT, size=10)
        cell.fill = PatternFill("solid", fgColor=TITRE_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for i, (idx, row) in enumerate(df.iterrows(), start=depart + 1):
        col = 1
        if index_titre:
            ws.cell(row=i, column=col, value=str(idx))
            col += 1
        for valeur in row:
            if pd.isna(valeur):
                col += 1
                continue
            if hasattr(valeur, "item"):
                valeur = valeur.item()
            ws.cell(row=i, column=col,
                    value=valeur if isinstance(valeur, (int, float, str))
                    else str(valeur))
            col += 1
    return depart + len(df) + 1


def _largeurs(ws, largeurs: dict[str, int]) -> None:
    for lettre, largeur in largeurs.items():
        ws.column_dimensions[lettre].width = largeur


def _paragraphe(ws, texte: str, ligne: int, largeur: int = 108, font=None) -> int:
    """Écrit un paragraphe sur plusieurs lignes de cellules.

    Une cellule fusionnée avec renvoi à la ligne paraît correcte dans Excel et
    déborde à l'impression sous LibreOffice, faute de hauteur de ligne calculée.
    Découper le texte à la main donne le même rendu partout.
    """
    import textwrap

    from openpyxl.styles import Font

    font = font or Font(size=10)
    for fragment in textwrap.wrap(texte, width=largeur) or [""]:
        c = ws.cell(row=ligne, column=1, value=fragment)
        c.font = font
        ligne += 1
    return ligne


# Feuilles trop larges pour tenir sur une page : les forcer à une largeur les
# rend illisibles à l'impression. Elles s'impriment sur plusieurs pages, comme
# le faisait le suivi tenu à la main.
LARGES = {"Détail par article", "Synthèse thématique", "Détail"}


def _mise_en_page(ws, paysage: bool = True) -> None:
    """Impression : orientation adaptée, ajustement seulement si c'est lisible."""
    ws.page_setup.orientation = "landscape" if paysage else "portrait"
    if ws.title in LARGES:
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.print_title_rows = "1:4"      # les en-têtes se répètent
        return
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _titre_feuille(ws, titre: str, sous_titre: str = "") -> int:
    """Pose un titre en A1 et renvoie la première ligne libre."""
    from openpyxl.styles import Font

    ws["A1"] = titre
    ws["A1"].font = Font(bold=True, size=13)
    if sous_titre:
        ws["A2"] = sous_titre
        ws["A2"].font = Font(size=9, color="52514E")
        return 4
    return 3


# ---------------------------------------------------------------------------
# Feuilles
# ---------------------------------------------------------------------------

def _feuille_lecture(wb, titre: str, dossier: str, matrix: pd.DataFrame,
                     detail: pd.DataFrame, methodologie: str,
                     palette: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.active
    ws.title = "Lecture"
    ligne = _titre_feuille(
        ws, titre,
        f"{dossier} · établi le {date.today().strftime('%d/%m/%Y')}")

    chiffres = [
        ("Articles couverts", len(matrix.index)),
        ("États membres", len(matrix.columns)),
        ("Contributions analysées", len(detail)),
        ("Cases renseignées", int(matrix.notna().sum().sum())),
        ("Taux de couverture",
         f"{matrix.notna().sum().sum() / max(matrix.size, 1):.0%}"),
    ]
    ws.cell(row=ligne, column=1, value="Chiffres clés").font = Font(bold=True)
    ligne += 1
    for intitule, valeur in chiffres:
        ws.cell(row=ligne, column=1, value=intitule)
        ws.cell(row=ligne, column=2, value=valeur)
        ligne += 1

    ligne += 1
    ws.cell(row=ligne, column=1, value="Légende de la matrice").font = Font(bold=True)
    ligne += 1
    pal = PALETTES[palette]
    legende = [
        (SYMBOLES[2], "aligné sur la position française", pal["aligne"]),
        (SYMBOLES[1], "alignement partiel, réserve ou condition", pal["partiel"]),
        (SYMBOLES[0], "divergent", pal["oppose"]),
        ("", "case vide : aucune contribution écrite sur cet article", pal["absent"]),
    ]
    for symbole, texte, couleur in legende:
        c = ws.cell(row=ligne, column=1, value=symbole)
        c.fill = PatternFill("solid", fgColor=_hex(couleur))
        c.alignment = Alignment(horizontal="center")
        ws.cell(row=ligne, column=2, value=texte)
        ligne += 1

    if methodologie:
        ligne += 1
        ws.cell(row=ligne, column=1, value="Méthode").font = Font(bold=True)
        ligne += 1
        ligne = _paragraphe(ws, methodologie, ligne)

    ligne += 1
    ligne = _paragraphe(
        ws,
        "Document produit automatiquement. Chaque position est adossée à une "
        "citation vérifiée dans le document source ; elle reste à valider par "
        "l'analyste avant tout usage en négociation.",
        ligne, font=Font(size=9, italic=True, color="898781"))

    _largeurs(ws, {"A": 30, "B": 46, "C": 14, "D": 14, "E": 14, "F": 14})


def _feuille_matrice(wb, matrix: pd.DataFrame, labels: dict, palette: str):
    """Matrice en damier, colorée et symbolisée, avec volets figés."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("Matrice")
    ligne = _titre_feuille(
        ws, "Positions article par article",
        "2 = aligné · 1 = partiel · 0 = divergent · vide = pas de contribution")

    pal = PALETTES[palette]
    couleurs = {2: _hex(pal["aligne"]), 1: _hex(pal["partiel"]),
                0: _hex(pal["oppose"])}

    ws.cell(row=ligne, column=1, value="Article").font = Font(bold=True, size=10)
    ws.cell(row=ligne, column=1).fill = PatternFill("solid", fgColor=TITRE_FILL)
    for j, col in enumerate(matrix.columns, start=2):
        c = ws.cell(row=ligne, column=j, value=str(col))
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=TITRE_FILL)
        c.alignment = Alignment(horizontal="center")

    entete = ligne
    for i, (idx, row) in enumerate(matrix.iterrows(), start=ligne + 1):
        ws.cell(row=i, column=1, value=labels.get(idx, str(idx)))
        for j, col in enumerate(matrix.columns, start=2):
            v = row[col]
            c = ws.cell(row=i, column=j)
            if pd.isna(v):
                c.fill = PatternFill("solid", fgColor=_hex(pal["absent"]))
                continue
            v = int(v)
            # La valeur numérique est conservée dans la cellule — elle sert aux
            # tris, aux formules et aux graphiques. Le symbole va en commentaire
            # de format, pas à la place du chiffre.
            c.value = v
            c.number_format = f'"{SYMBOLES[v]}"'
            c.fill = PatternFill("solid", fgColor=couleurs[v])
            c.alignment = Alignment(horizontal="center")

    fin = entete + len(matrix.index)
    ws.freeze_panes = ws.cell(row=entete + 1, column=2)
    ws.auto_filter.ref = (f"A{entete}:"
                          f"{ws.cell(row=entete, column=len(matrix.columns) + 1).column_letter}"
                          f"{fin}")
    _largeurs(ws, {"A": 34})
    for j in range(2, len(matrix.columns) + 2):
        ws.column_dimensions[ws.cell(row=entete, column=j).column_letter].width = 5.5
    return ws


def _feuille_classement(wb, ranking: pd.DataFrame):
    """Classement des États membres + histogramme natif, lié aux cellules."""
    from openpyxl.chart import BarChart, Reference

    ws = wb.create_sheet("Classement EM")
    ligne = _titre_feuille(
        ws, "Proximité avec la position française",
        "Score moyen sur les articles où l'État membre s'est exprimé "
        "(2 = aligné, 0 = divergent)")
    if ranking.empty:
        return ws

    depart = ligne
    _ecrire_df(ws, ranking, index_titre="État membre", depart=depart)
    fin = depart + len(ranking)

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Score moyen d'alignement"
    chart.y_axis.title = "Score moyen"
    chart.x_axis.title = "État membre"
    chart.height, chart.width = max(8, 0.45 * len(ranking) + 3), 16
    donnees = Reference(ws, min_col=2, min_row=depart, max_row=fin)
    categories = Reference(ws, min_col=1, min_row=depart + 1, max_row=fin)
    chart.add_data(donnees, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend = None
    ws.add_chart(chart, f"J{depart}")

    _largeurs(ws, {"A": 16, "B": 14, "C": 18, "D": 14, "E": 12, "F": 14})
    ws.freeze_panes = ws.cell(row=depart + 1, column=1)
    return ws


def _feuille_clivants(wb, contentious: pd.DataFrame, labels: dict):
    """Articles clivants + nuage de points natif (dispersion × score)."""
    from openpyxl.chart import Reference, ScatterChart, Series

    ws = wb.create_sheet("Articles clivants")
    ligne = _titre_feuille(
        ws, "Articles par niveau de consensus",
        "Dispersion élevée = les États membres se divisent ; score bas = "
        "opposition large")
    if contentious.empty:
        return ws

    table = contentious.copy()
    table.index = [labels.get(i, str(i)) for i in table.index]
    depart = ligne
    _ecrire_df(ws, table, index_titre="Article", depart=depart)
    fin = depart + len(table)

    colonnes = list(table.columns)
    if "Score moyen" in colonnes and "Dispersion" in colonnes:
        col_score = colonnes.index("Score moyen") + 2
        col_disp = colonnes.index("Dispersion") + 2
        chart = ScatterChart()
        chart.title = "Score moyen × dispersion"
        chart.x_axis.title = "Score moyen"
        chart.y_axis.title = "Dispersion"
        chart.height, chart.width = 11, 16
        xs = Reference(ws, min_col=col_score, min_row=depart + 1, max_row=fin)
        ys = Reference(ws, min_col=col_disp, min_row=depart, max_row=fin)
        serie = Series(ys, xs, title_from_data=True)
        serie.marker.symbol = "circle"
        serie.graphicalProperties.line.noFill = True   # nuage, pas courbe
        chart.series.append(serie)
        chart.legend = None
        ws.add_chart(chart, f"J{depart}")

    _largeurs(ws, {"A": 34, "B": 14, "C": 14, "D": 14, "E": 14})
    ws.freeze_panes = ws.cell(row=depart + 1, column=1)
    return ws


def _feuille_couverture(wb, detail: pd.DataFrame):
    """Engagement de chaque État membre : combien d'articles, combien de fois."""
    from openpyxl.chart import BarChart, Reference

    ws = wb.create_sheet("Couverture")
    ligne = _titre_feuille(
        ws, "Engagement des États membres",
        "Nombre de contributions écrites et d'articles couverts")
    if detail.empty:
        return ws

    table = (detail.groupby("ms_code")
             .agg(Contributions=("text", "size"),
                  Articles=("section_label", "nunique"))
             .sort_values("Contributions", ascending=False))
    depart = ligne
    _ecrire_df(ws, table, index_titre="État membre", depart=depart)
    fin = depart + len(table)

    chart = BarChart()
    chart.type = "col"
    chart.title = "Contributions écrites par État membre"
    chart.height, chart.width = 10, 20
    donnees = Reference(ws, min_col=2, max_col=3, min_row=depart, max_row=fin)
    categories = Reference(ws, min_col=1, min_row=depart + 1, max_row=fin)
    chart.add_data(donnees, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, f"F{depart}")

    _largeurs(ws, {"A": 16, "B": 16, "C": 14})
    ws.freeze_panes = ws.cell(row=depart + 1, column=1)
    return ws


# Colonnes du détail, dans l'ordre de lecture d'un analyste : qui, sur quoi,
# quelle position, pourquoi, où le vérifier.
COLONNES_DETAIL = [
    ("ms_code", "État membre", 12),
    ("section_label", "Article", 26),
    ("disposition", "Disposition visée", 18),
    ("themes_fr", "Thèmes", 26),
    ("stance_fr", "Position", 14),
    ("summary_fr", "Ce que dit l'État membre", 62),
    ("evidence", "Citation vérifiée dans le texte source", 62),
    ("scrutiny_fr", "Réserve d'examen", 14),
    ("deletion_fr", "Demande de suppression", 16),
    ("method", "Méthode", 12),
    ("confidence", "Confiance", 10),
    ("page_start", "Page", 8),
    ("text", "Texte intégral de la contribution", 90),
]


def _feuille_detail(wb, detail: pd.DataFrame):
    """Une ligne par contribution, avec le texte — le cœur du classeur.

    C'est la feuille qui permet de contester une case de la matrice sans
    rouvrir le PDF : le résumé, la citation qui le fonde et la page y figurent
    côte à côte.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("Détail")
    ligne = _titre_feuille(
        ws, "Détail des contributions",
        "Une ligne par contribution écrite. La citation est vérifiée "
        "littéralement dans le document source.")
    if detail.empty:
        return ws

    presentes = [(c, t, w) for c, t, w in COLONNES_DETAIL if c in detail.columns]
    for j, (_, titre, _) in enumerate(presentes, start=1):
        c = ws.cell(row=ligne, column=j, value=titre)
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=TITRE_FILL)
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for i, (_, row) in enumerate(detail.iterrows(), start=ligne + 1):
        for j, (col, _, _) in enumerate(presentes, start=1):
            v = row[col]
            if pd.isna(v):
                continue
            if hasattr(v, "item"):
                v = v.item()
            c = ws.cell(row=i, column=j,
                        value=v if isinstance(v, (int, float)) else str(v))
            c.alignment = Alignment(wrap_text=True, vertical="top")

    fin = ligne + len(detail)
    ws.freeze_panes = ws.cell(row=ligne + 1, column=3)
    ws.auto_filter.ref = (f"A{ligne}:"
                          f"{ws.cell(row=ligne, column=len(presentes)).column_letter}"
                          f"{fin}")
    for j, (_, _, largeur) in enumerate(presentes, start=1):
        ws.column_dimensions[ws.cell(row=ligne, column=j).column_letter].width = largeur
    return ws


def _feuille_figures(wb, images: list[tuple[str, bytes]]):
    """Les figures que le graphique natif ne rend pas — la matrice, surtout."""
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Font

    ws = wb.create_sheet("Figures")
    _titre_feuille(ws, "Figures",
                   "Images fixes. Les graphiques liés aux données sont sur les "
                   "feuilles « Classement EM », « Articles clivants » et "
                   "« Couverture ».")
    ligne = 4
    for legende, png in images:
        ws.cell(row=ligne, column=1, value=legende).font = Font(bold=True, size=11)
        ligne += 1
        try:
            img = Image(io.BytesIO(png))
            # Excel raisonne en pixels ; on plafonne pour que la figure tienne
            # dans une page à l'impression.
            ratio = min(1.0, 1100 / max(img.width, 1))
            img.width, img.height = int(img.width * ratio), int(img.height * ratio)
            ws.add_image(img, f"A{ligne}")
            ligne += int(img.height / 19) + 3
        except Exception:
            # Une image absente ne doit pas emporter le classeur entier.
            ws.cell(row=ligne, column=1, value="[figure non intégrée]")
            ligne += 2
    _largeurs(ws, {"A": 22})
    return ws


# ---------------------------------------------------------------------------
# Feuilles « à la main » — celles qui reprennent la forme du suivi existant
# ---------------------------------------------------------------------------
#
# Le suivi tenu à la main par la direction a une forme éprouvée : une ligne par
# article, une colonne par État membre, la France en première colonne comme
# référence, et la couleur de la case qui dit la compatibilité avec elle. Cette
# forme se lit d'un coup d'œil et se projette en réunion.
#
# L'outil la reproduit plutôt que d'en imposer une autre. La différence est
# ailleurs : ici chaque case est produite à partir d'une contribution
# réellement citée, et le classeur se régénère en trente secondes au lieu de
# se tenir à la main.

MS_NOMS = {
    "AT": "Autriche", "BE": "Belgique", "BG": "Bulgarie", "CY": "Chypre",
    "CZ": "Tchéquie", "DE": "Allemagne", "DK": "Danemark", "EE": "Estonie",
    "EL": "Grèce", "ES": "Espagne", "FI": "Finlande", "FR": "France",
    "HR": "Croatie", "HU": "Hongrie", "IE": "Irlande", "IT": "Italie",
    "LT": "Lituanie", "LU": "Luxembourg", "LV": "Lettonie", "MT": "Malte",
    "NL": "Pays-Bas", "PL": "Pologne", "PT": "Portugal", "RO": "Roumanie",
    "SE": "Suède", "SI": "Slovénie", "SK": "Slovaquie",
}

# Longueur d'une case : assez pour porter la position, pas au point de faire
# une page de texte dans une cellule.
MAX_CASE = 420


def _entete_ms(ws, ligne: int, colonnes: list[str], pal: dict,
               titre_premiere: str = "Article") -> None:
    """En-tête : intitulé, puis FR en première colonne, puis les autres EM."""
    from openpyxl.styles import Alignment, Font, PatternFill

    c = ws.cell(row=ligne, column=1, value=titre_premiere)
    c.font = Font(bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=_hex(pal["entete"]))
    c.alignment = Alignment(vertical="center", wrap_text=True)

    for j, code in enumerate(colonnes, start=2):
        cellule = ws.cell(row=ligne, column=j,
                          value=f"{code}\n{MS_NOMS.get(code, '')}".strip())
        cellule.font = Font(bold=True, size=10, color="FFFFFF")
        cellule.fill = PatternFill(
            "solid",
            fgColor=_hex(pal["entete"] if code == "FR" else pal["entete_ms"]))
        cellule.alignment = Alignment(horizontal="center", vertical="center",
                                      wrap_text=True)


def _case(ws, ligne: int, colonne: int, texte: str, couleur: str,
          gras: bool = False) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    c = ws.cell(row=ligne, column=colonne, value=texte[:MAX_CASE])
    c.fill = PatternFill("solid", fgColor=_hex(couleur))
    # Encre noire explicite : sur un fond vert ou rouge, la couleur de police
    # héritée du thème peut passer en blanc et devenir illisible à l'écran
    # comme à l'impression.
    c.font = Font(size=9, bold=gras, color="0B0B0B")
    c.alignment = Alignment(wrap_text=True, vertical="top")


def _texte_de_case(lignes: pd.DataFrame) -> str:
    """Ce qu'on écrit dans une case : la position, puis sa référence.

    Deux contributions du même État sur le même article sont réunies, la plus
    défavorable en tête — c'est celle qui décide de la couleur, elle doit être
    celle qu'on lit en premier.
    """
    if lignes.empty:
        return "—"
    rang = {"oppose": 0, "partiel": 1, "aligne": 2, "neutre": 3}
    ordonne = lignes.assign(
        _r=lignes["stance"].map(lambda s: rang.get(s, 4))).sort_values("_r")

    morceaux = []
    for _, row in ordonne.head(2).iterrows():
        resume = str(row.get("summary_fr") or "").strip()
        if not resume:
            resume = str(row.get("evidence") or "").strip()
        resume = resume.replace("[Classement lexical, non validé] ", "")
        reference = str(row.get("disposition") or "").strip()
        if not reference or reference == "—":
            page = row.get("page_start")
            reference = f"p. {int(page)}" if pd.notna(page) else ""
        morceaux.append(resume + (f"\n[{reference}]" if reference else ""))
    reste = len(ordonne) - 2
    if reste > 0:
        morceaux.append(f"(+ {reste} autre(s) contribution(s))")
    return "\n".join(morceaux)


def _feuille_position_fr(wb, fr_positions: dict, sujets: dict,
                         articles: list, pal: dict):
    """La position française de référence, article par article.

    Première feuille consultée en réunion : sans elle, « opposé » ne veut rien
    dire. Elle dit aussi, article par article, si la référence est un
    amendement écrit ou la règle du silence.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("Position FR")
    ligne = _titre_feuille(
        ws, "Position française de référence",
        "Sur un article que la France n'a pas amendé, la référence est le "
        "maintien du texte initial en l'état : ne pas amender vaut "
        "acceptation.")

    entetes = ["Article", "Intitulé", "Position française de référence",
               "Origine"]
    for j, titre in enumerate(entetes, start=1):
        c = ws.cell(row=ligne, column=j, value=titre)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_hex(pal["entete"]))
        c.alignment = Alignment(vertical="center", wrap_text=True)

    depart = ligne
    for i, article in enumerate(articles, start=ligne + 1):
        texte = (fr_positions.get(article) or "").strip()
        explicite = bool(texte)
        _case(ws, i, 1, str(article), pal["reference"], gras=True)
        _case(ws, i, 2, str(sujets.get(article, "") or "—"), "FFFFFF")
        _case(ws, i, 3,
              texte or "Aucun amendement français : maintien du texte en l'état.",
              "FFFFFF")
        _case(ws, i, 4,
              "amendement écrit" if explicite else "règle du silence",
              pal["reference"] if explicite else pal["absent"])

    ws.freeze_panes = ws.cell(row=depart + 1, column=2)
    _largeurs(ws, {"A": 18, "B": 40, "C": 80, "D": 18})
    return ws


def _feuille_themes(wb, detail: pd.DataFrame, pal: dict):
    """Synthèse thématique : les enjeux qui traversent plusieurs articles.

    Un dossier ne se négocie pas article par article mais enjeu par enjeu —
    comitologie, rôle des États membres, périmètre. Les thèmes sont ceux
    relevés dans les contributions elles-mêmes ; la couleur d'une case est la
    position la plus défavorable exprimée par cet État membre sur ce thème.
    """
    import json

    ws = wb.create_sheet("Synthèse thématique")
    ligne = _titre_feuille(
        ws, "Synthèse thématique",
        "Position de chaque État membre par enjeu, tous articles confondus. "
        "Couleur = position la plus défavorable exprimée sur cet enjeu.")

    if detail is None or detail.empty or "themes" not in detail.columns:
        _paragraphe(ws, "Aucun thème relevé : cette feuille se remplit quand "
                        "l'analyse est faite avec le modèle, qui relève les "
                        "thèmes de chaque contribution.", ligne)
        _largeurs(ws, {"A": 44})
        return ws

    lignes = []
    for _, row in detail.iterrows():
        try:
            themes = json.loads(row.get("themes") or "[]")
        except (TypeError, ValueError):
            themes = []
        for t in themes:
            lignes.append({**row.to_dict(), "theme": str(t).strip().capitalize()})
    if not lignes:
        _paragraphe(ws, "Aucun thème relevé dans les contributions analysées.",
                    ligne)
        _largeurs(ws, {"A": 44})
        return ws

    par_theme = pd.DataFrame(lignes)
    # Les thèmes les plus portés d'abord : ce sont les enjeux du dossier.
    ordre_themes = list(par_theme["theme"].value_counts().head(20).index)
    colonnes = _colonnes_ms(par_theme)

    _entete_ms(ws, ligne, colonnes, pal, titre_premiere="Enjeu")
    couleurs = _couleurs_stance(pal)

    for i, theme in enumerate(ordre_themes, start=ligne + 1):
        bloc = par_theme[par_theme["theme"] == theme]
        _case(ws, i, 1, f"{theme}\n({len(bloc)} contributions)",
              pal["reference"], gras=True)
        for j, code in enumerate(colonnes, start=2):
            cellules = bloc[bloc["ms_code"] == code]
            couleur = (pal["reference"] if code == "FR"
                       else couleurs.get(_stance_dominante(cellules), pal["absent"]))
            _case(ws, i, j, _texte_de_case(cellules), couleur,
                  gras=(code == "FR"))
        ws.row_dimensions[i].height = 74

    ws.freeze_panes = ws.cell(row=ligne + 1, column=2)
    _largeurs(ws, {"A": 34})
    for j in range(2, len(colonnes) + 2):
        ws.column_dimensions[
            ws.cell(row=ligne, column=j).column_letter].width = 34
    return ws


def _colonnes_ms(detail: pd.DataFrame) -> list[str]:
    """États membres présents, la France toujours en première colonne."""
    codes = sorted({str(c) for c in detail.get("ms_code", []) if str(c).strip()})
    autres = [c for c in codes if c != "FR"]
    return (["FR"] if "FR" in codes else []) + autres


def _couleurs_stance(pal: dict) -> dict:
    return {"aligne": pal["aligne"], "partiel": pal["partiel"],
            "oppose": pal["oppose"], "neutre": pal["absent"]}


def _stance_dominante(lignes: pd.DataFrame) -> str:
    """La position la plus défavorable exprimée — celle qui donne la couleur."""
    if lignes is None or lignes.empty:
        return ""
    for stance in ("oppose", "partiel", "aligne", "neutre"):
        if (lignes["stance"] == stance).any():
            return stance
    return ""


def _feuille_detail_par_article(wb, matrix: pd.DataFrame, detail: pd.DataFrame,
                                fr_positions: dict, sujets: dict, labels: dict,
                                pal: dict):
    """Le tableau de suivi lui-même : articles × États membres, en couleur.

    C'est la feuille qui remplace le suivi tenu à la main. Colonne A :
    l'article et son sujet. Colonne B : la position française, qui donne son
    sens à tout le reste — sans elle, « opposé » ne se lit pas.
    """
    ws = wb.create_sheet("Détail par article")
    ligne = _titre_feuille(
        ws, "Détail article par article",
        "Position de chaque État membre, couleur = compatibilité avec la "
        "position française. La colonne FR porte la référence.")

    if detail is None or detail.empty:
        _paragraphe(ws, "Aucune contribution analysée.", ligne)
        return ws

    colonnes = [c for c in _colonnes_ms(detail) if c != "FR"]
    _entete_ms(ws, ligne, ["FR"] + colonnes, pal)
    couleurs = _couleurs_stance(pal)

    articles = list(matrix.index) if not matrix.empty else list(
        dict.fromkeys(detail["section_label"]))

    for i, article in enumerate(articles, start=ligne + 1):
        intitule = str(labels.get(article, article))
        sujet = str(sujets.get(article, "") or "").strip()
        _case(ws, i, 1, intitule + (f"\n{sujet}" if sujet else ""),
              pal["reference"], gras=True)

        reference = (fr_positions.get(article) or "").strip()
        _case(ws, i, 2,
              reference or "Aucun amendement français : maintien du texte en "
                           "l'état.\n[règle du silence]",
              pal["reference"], gras=True)

        bloc = detail[detail["section_label"] == article]
        for j, code in enumerate(colonnes, start=3):
            cellules = bloc[bloc["ms_code"] == code]
            couleur = couleurs.get(_stance_dominante(cellules), pal["absent"])
            _case(ws, i, j, _texte_de_case(cellules), couleur)
        ws.row_dimensions[i].height = 74

    fin = ligne + len(articles)
    ws.freeze_panes = ws.cell(row=ligne + 1, column=3)
    ws.auto_filter.ref = (
        f"A{ligne}:"
        f"{ws.cell(row=ligne, column=len(colonnes) + 2).column_letter}{fin}")
    _largeurs(ws, {"A": 30, "B": 38})
    for j in range(3, len(colonnes) + 3):
        ws.column_dimensions[
            ws.cell(row=ligne, column=j).column_letter].width = 34
    return ws


# ---------------------------------------------------------------------------
# Point d'entrée
# ---------------------------------------------------------------------------

def build_xlsx(
    titre: str,
    dossier: str,
    matrix: pd.DataFrame,
    ranking: pd.DataFrame,
    contentious: pd.DataFrame,
    detail: pd.DataFrame,
    images: list[tuple[str, bytes]] | None = None,
    labels: dict | None = None,
    palette: str | None = None,
    methodologie: str = "",
    fr_positions: dict | None = None,
    sujets: dict | None = None,
) -> bytes:
    """Produit le classeur complet.

    `detail` est la table des contributions analysées ; les colonnes absentes
    sont simplement ignorées, ce qui permet d'appeler la fonction depuis un
    contexte où l'analyse n'a pas encore tourné.
    """
    from openpyxl import Workbook

    labels = labels or {}
    fr_positions = fr_positions or {}
    sujets = sujets or {}
    palette = palette or actif()
    pal = PALETTES[palette]
    wb = Workbook()

    detail = _preparer_detail(detail)

    # Ordre des feuilles : celui du suivi tenu à la main. On lit d'abord la
    # position française, puis les enjeux, puis le détail — et seulement
    # ensuite les feuilles de calcul, qui servent aux graphiques.
    _feuille_lecture(wb, titre, dossier, matrix, detail, methodologie, palette)
    _feuille_position_fr(wb, fr_positions, sujets,
                         list(matrix.index) if not matrix.empty else [], pal)
    _feuille_themes(wb, detail, pal)
    _feuille_detail_par_article(wb, matrix, detail, fr_positions, sujets,
                                labels, pal)
    _feuille_matrice(wb, matrix, labels, palette)
    _feuille_classement(wb, ranking)
    _feuille_clivants(wb, contentious, labels)
    _feuille_couverture(wb, detail)
    _feuille_detail(wb, detail)
    if images:
        _feuille_figures(wb, images)

    for ws in wb.worksheets:
        _mise_en_page(ws, paysage=ws.title != "Lecture")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _preparer_detail(detail: pd.DataFrame) -> pd.DataFrame:
    """Traduit les colonnes techniques en libellés lisibles dans Excel."""
    import json

    from .schemas import STANCE_LABEL_FR

    if detail is None or detail.empty:
        return pd.DataFrame()

    d = detail.copy()
    if "stance" in d.columns:
        d["stance_fr"] = d["stance"].map(
            lambda s: STANCE_LABEL_FR.get(s, "—") if pd.notna(s) else "—")
    for col, sortie in (("scrutiny", "scrutiny_fr"), ("deletion", "deletion_fr")):
        if col in d.columns:
            d[sortie] = d[col].map(lambda v: "oui" if v else "")
    if "themes" in d.columns:
        def _themes(v):
            try:
                return ", ".join(json.loads(v)) if v else ""
            except (TypeError, ValueError):
                return ""
        d["themes_fr"] = d["themes"].map(_themes)
    # Disposition réellement visée par la délégation : « § 2 »,
    # « considérant 12 ». Sans elle, dix lignes « Article 5 » ne se
    # distinguent pas les unes des autres dans un tableau filtrable.
    if "text" in d.columns:
        from .labels import reference_dans_texte

        d["disposition"] = d["text"].map(
            lambda t: reference_dans_texte(t or "") or "—")
    if "section_order" in d.columns:
        d = d.sort_values(["section_order", "ms_code"])
    return d

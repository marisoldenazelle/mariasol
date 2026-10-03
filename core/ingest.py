"""
Ingestion universelle des documents de négociation.

Trois familles de documents, détectées automatiquement mais toujours
rectifiables à la main lors de l'import :

  wk_table    tableau de commentaires consolidés du Conseil (2 ou 3 colonnes)
  legal_text  texte réglementaire : proposition, compromis de présidence,
              règlement publié — découpé article par article
  free_text   non-paper, note de position, compte rendu de réunion —
              découpé par paragraphe, avec attribution d'État membre
              quand le document en désigne un

Tout est déterministe. Aucun appel réseau à ce stade : un document mal
découpé doit se voir et se corriger, pas se deviner.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import pdfplumber

from .wk_parser import MS_CODES, MS_NAMES, merge_contributions, parse_wk_pdf

DOC_KINDS = {
    "wk_table": "Commentaires consolidés du Conseil (tableau, 2 ou 3 colonnes)",
    "legal_text": "Texte réglementaire (proposition, compromis, règlement)",
    "free_text": "Texte libre (non-paper, note de position, compte rendu)",
}


@dataclass
class Segment:
    """Unité indexable commune à tous les types de documents."""

    section_id: str
    section_label: str
    section_kind: str            # article | recital | annex | general | paragraphe
    text: str
    page_start: int
    page_end: int
    ms_code: str = ""            # vide si le segment n'est pas attribuable
    contribution_kind: str = ""  # comments | drafting | ""
    order: int = 0


# ---------------------------------------------------------------------------
# Lecture brute
# ---------------------------------------------------------------------------

def read_pages(path: str | Path) -> list[str]:
    """Renvoie le texte de chaque page (PDF) ou du document entier (docx)."""
    path = Path(path)
    if path.suffix.lower() in (".docx", ".dotx"):
        from docx import Document

        doc = Document(str(path))
        blocks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                blocks.append(" | ".join(c.text.strip() for c in row.cells))
        # séparation par ligne vide : le découpage par paragraphe en dépend
        return ["\n\n".join(b.strip() for b in blocks if b.strip())]

    with pdfplumber.open(str(path)) as pdf:
        return [_sans_pied_de_page(p.extract_text() or "") for p in pdf.pages]


# Pied de page du Journal officiel et des documents COM : le code de langue
# répété, puis le numéro de page. « FR FR » suivi de « 23 ».
_PIED_DE_PAGE = re.compile(
    r"\n[ \t]*(?:[A-Z]{2}[ \t]*){1,3}\n[ \t]*\d{1,4}[ \t]*$")


def _sans_pied_de_page(page: str) -> str:
    """Retire le pied de page courant d'un document européen.

    Le PDF d'un document COM porte en bas de chaque page « FR FR » et le
    numéro. L'extraction le colle à la dernière phrase, et cette phrase est
    souvent une instruction de modification : l'écran affichait « au
    paragraphe 2, les points suivants sont ajoutés: FR FR 23 ». Le pied doit
    donc partir à la lecture, pas dans chaque module qui lit ensuite le texte.

    On n'enlève que ce motif-là, en fin de page, et rien d'autre : un texte
    qui se terminerait vraiment par un nombre seul est plus rare qu'un pied
    de page, mais la règle exige les deux lignes.
    """
    return _PIED_DE_PAGE.sub("", page or "").rstrip()


# ---------------------------------------------------------------------------
# Détection du type
# ---------------------------------------------------------------------------

_TWO_COL_HEADER = re.compile(r"commission proposal", re.IGNORECASE)
# « Article 5 », « Article 5 bis », « Article premier » (article 1 en français),
# seuls sur leur ligne ou suivis de leur intitulé.
_ARTICLE_HEAD = re.compile(
    r"^\s*Article\s+(?:premier|\d+\s*(?:bis|ter|quater)?[a-z]?)\s*$",
    re.IGNORECASE | re.MULTILINE)


def detect_kind(path: str | Path, pages: list[str] | None = None) -> str:
    """Propose un type de document. L'utilisateur peut toujours le corriger."""
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        try:
            with pdfplumber.open(str(path)) as pdf:
                sample = pdf.pages[: min(6, len(pdf.pages))]
                # Deux formats coexistent : « proposition | commentaires » et
                # « proposition | compromis de présidence | commentaires ».
                tabular = sum(
                    1 for p in sample
                    for t in p.find_tables()
                    if len(t.extract()[0] or []) in (2, 3)
                )
                if tabular >= 2:
                    return "wk_table"
        except Exception:
            pass

    pages = pages if pages is not None else read_pages(path)
    if _TWO_COL_HEADER.search("\n".join(pages[:12])):
        return "wk_table"
    # On balaie tout le document, pas ses premières pages : un règlement
    # européen ouvre sur cent à deux cents considérants, et le premier article
    # n'arrive qu'ensuite. Regarder les douze premières pages classait le RGPD
    # et l'AI Act en « texte libre ».
    joined = "\n".join(pages)
    if len(_ARTICLE_HEAD.findall(joined)) >= 3:
        return "legal_text"
    return "free_text"


# ---------------------------------------------------------------------------
# Découpage — texte réglementaire
# ---------------------------------------------------------------------------

_ART_START = re.compile(
    r"(?:^|\n)\s*(?P<head>Article\s+(?P<num>premier|\d+\s*(?:bis|ter|quater)?[a-z]?))"
    r"\s*(?:\n|$)",
    re.IGNORECASE,
)
_REC_START = re.compile(r"(?:^|\n)\s*\((?P<num>\d+[a-z]?)\)\s+(?=[A-Z])")

# Début de l'acte lui-même. Ce qui précède, dans une proposition de la
# Commission, est l'exposé des motifs et la fiche financière législative :
# cent pages qui ne sont pas du droit, et qui atterrissaient dans le même
# segment que les considérants. Elles remontaient alors en recherche, et un
# résumé pouvait présenter un tableau budgétaire comme une disposition.
_DEBUT_ACTE = re.compile(
    r"(?:LE\s+PARLEMENT\s+EUROP[ÉE]EN\s+ET\s+LE\s+CONSEIL|"
    r"LA\s+COMMISSION\s+EUROP[ÉE]ENNE|LE\s+CONSEIL\s+DE\s+L[’']UNION|"
    r"vu\s+le\s+trait[ée]\s+sur\s+le\s+fonctionnement)",
    re.IGNORECASE)

# Formule de clôture. Ce qui suit — annexes financières, tableaux de
# correspondance — n'appartient plus au dernier article, qui l'avalait :
# l'article 11 de l'omnibus faisait quatre-vingt-dix mille caractères.
_FIN_ACTE = re.compile(
    r"(?m)^\s*(?:Fait\s+à\s+\w+|Par\s+le\s+Parlement\s+europ[ée]en)\b")


def split_legal_text(pages: list[str]) -> list[Segment]:
    """Découpe un texte réglementaire à la granularité de l'article.

    Le préambule (considérants) est conservé comme une section propre : dans
    les négociations, les considérants portent une part réelle du compromis.
    """
    offsets: list[tuple[int, int]] = []   # (position dans le texte joint, page)
    buf: list[str] = []
    cursor = 0
    for i, text in enumerate(pages, start=1):
        offsets.append((cursor, i))
        buf.append(text)
        cursor += len(text) + 1
    joined = "\n".join(buf)

    def page_of(pos: int) -> int:
        page = 1
        for start, p in offsets:
            if start <= pos:
                page = p
            else:
                break
        return page

    # on repère la position du titre lui-même, pas du saut de ligne qui le
    # précède : sinon un article ouvrant une page est attribué à la précédente
    marks = [(m.start("head"), m.group("num"), m.group("head"))
             for m in _ART_START.finditer(joined)]
    segments: list[Segment] = []

    if marks and marks[0][0] > 60:
        head = joined[: marks[0][0]].strip()
        if head:
            # L'exposé des motifs est séparé des considérants quand les deux
            # coexistent : ce sont deux natures de texte, et une seule des
            # deux se négocie.
            coupe = _DEBUT_ACTE.search(head)
            expose = head[:coupe.start()].strip() if coupe else ""
            preambule = head[coupe.start():].strip() if coupe else head
            if len(expose) > 800:
                segments.append(Segment(
                    "expose", "Exposé des motifs et fiche financière",
                    "expose", expose, 1, page_of(len(expose)), order=-2))
            else:
                preambule = head
            if preambule:
                segments.append(Segment(
                    "preambule", "Préambule et considérants", "recital",
                    preambule, page_of(len(head) - len(preambule)),
                    page_of(marks[0][0]), order=-1))

    cloture = _FIN_ACTE.search(joined)
    cloture_pos = cloture.start() if cloture else None

    for i, (pos, num, head) in enumerate(marks):
        num = re.sub(r"\s+", "_", num.strip().lower())
        if num == "premier":
            num = "1"          # « Article premier » est l'article 1
        end = marks[i + 1][0] if i + 1 < len(marks) else len(joined)
        # La formule de clôture ferme l'article qui la contient : ce qui suit
        # est signature, annexes et fiche financière. Sans cette coupe, le
        # dernier article de l'acte avalait quatre-vingt-dix mille caractères.
        if cloture_pos is not None and pos < cloture_pos < end:
            end = cloture_pos
        body = joined[pos:end].strip()
        if len(body) < 20:
            continue
        segments.append(Segment(
            section_id=f"art_{num}",
            section_label=f"Article {num.replace('_', ' ')}",
            section_kind="article",
            text=body,
            page_start=page_of(pos),
            page_end=page_of(max(end - 1, pos)),
            order=i,
        ))

    # Ce qui suit la formule de clôture n'est pas perdu : signature, annexes,
    # tableaux de correspondance, fiche financière. C'est une section à part,
    # nommée comme telle plutôt que rattachée au dernier article.
    if cloture_pos is not None and len(joined) - cloture_pos > 800:
        segments.append(Segment(
            "annexes", "Annexes et pièces jointes", "annexe",
            joined[cloture_pos:].strip(), page_of(cloture_pos),
            page_of(len(joined) - 1), order=9_000))

    if not segments:  # document sans structure d'article reconnaissable
        return split_free_text(pages)
    segments = _decouper_par_modification(pages, segments)
    return _subdiviser_articles(segments)


def _decouper_par_modification(pages: list[str],
                               segments: list[Segment]) -> list[Segment]:
    """Sur un acte modificatif, découpe le dispositif instruction par instruction.

    L'article premier d'un omnibus fait trente mille caractères et énumère
    quarante modifications de textes différents. Le garder d'un bloc rendait
    trois choses mauvaises : la recherche ramenait ce pavé entier, la
    comparaison de deux versions affichait un diff illisible, et le fil du
    texte ne proposait que « Article 1 ».

    Découpé, chaque instruction devient une section à part entière, dont
    l'identifiant porte le texte cible et l'article visé — `mod_32016R0679_
    art_33_8`. Deux versions du même omnibus s'apparient alors instruction par
    instruction : c'est ce qui permet de voir ce qu'un compromis de présidence
    change à la proposition de la Commission, article visé par article visé.

    Si la lecture échoue, on rend le découpage ordinaire : mieux vaut un pavé
    qu'un document tronqué.
    """
    from .eurlex import NOMS_USAGE
    from .modificatif import entetes_dispositif, est_acte_modificatif, lire

    joint = "\n".join(pages)
    if not est_acte_modificatif(joint):
        return segments

    try:
        lecture = lire(pages)
    except Exception:
        return segments
    if not lecture.blocs:
        return segments

    conserves = [s for s in segments if s.section_kind != "article"]
    entetes = set(entetes_dispositif(pages))
    modifiants = {("article 1" if "premier" in b.article_omnibus.lower()
                   else b.article_omnibus.lower()) for b in lecture.blocs}
    for seg in segments:
        if seg.section_kind != "article":
            continue
        label = re.sub(r"\s*\(\d+/\d+\)\s*$", "",
                       seg.section_label).strip().lower()
        # Un article que l'omnibus INSÈRE — « Article 32 bis » — n'est pas un
        # article de l'omnibus : c'est du texte cité, déjà porté par
        # l'instruction qui l'insère. Le garder en double faisait remonter le
        # même passage deux fois en recherche, et peuplait le fil du texte
        # d'articles qui n'existent pas encore.
        if label not in entetes:
            continue
        # Les articles qui ne modifient rien — entrée en vigueur, dispositions
        # finales — restent des articles ordinaires.
        if label not in modifiants:
            conserves.append(seg)

    ordre = 0
    for bloc in lecture.blocs:
        court = NOMS_USAGE.get(bloc.celex) or bloc.acte_cible
        for mod in bloc.modifications:
            ordre += 1
            cle = re.sub(r"[^a-z0-9]+", "_",
                         f"{bloc.celex or bloc.acte_cible}_{mod.section_id}_"
                         f"{mod.numero}".lower()).strip("_")
            corps = mod.instruction
            if mod.texte_nouveau:
                corps += f"\n« {mod.texte_nouveau} »"
            conserves.append(Segment(
                section_id=f"mod_{cle}",
                section_label=(f"{bloc.article_omnibus} · {mod.numero} · "
                               f"{court}"
                               + (f", {mod.article_cible}"
                                  if mod.article_cible else "")),
                section_kind="modification",
                text=corps,
                page_start=mod.page,
                page_end=mod.page,
                order=ordre,
            ))
    return conserves


# ---------------------------------------------------------------------------
# Découpage — texte libre
# ---------------------------------------------------------------------------

# « Non-paper by France », « FR comments », « Comments from the Netherlands »
_MS_HINT = re.compile(
    r"\b(?P<code>" + "|".join(MS_CODES) + r")\b|"
    + "|".join(rf"\b{re.escape(n)}\b" for n in MS_NAMES.values()),
    re.IGNORECASE,
)
_EN_NAMES = {
    "austria": "AT", "belgium": "BE", "bulgaria": "BG", "croatia": "HR",
    "cyprus": "CY", "czechia": "CZ", "czech republic": "CZ", "denmark": "DK",
    "estonia": "EE", "finland": "FI", "france": "FR", "germany": "DE",
    "greece": "EL", "hungary": "HU", "ireland": "IE", "italy": "IT",
    "latvia": "LV", "lithuania": "LT", "luxembourg": "LU", "malta": "MT",
    "netherlands": "NL", "poland": "PL", "portugal": "PT", "romania": "RO",
    "slovakia": "SK", "slovenia": "SI", "spain": "ES", "sweden": "SE",
}


def guess_author(pages: list[str]) -> str:
    """Devine l'État membre auteur d'un non-paper à partir de son en-tête."""
    head = "\n".join(pages[:2])[:1500].lower()
    for name, code in _EN_NAMES.items():
        if re.search(rf"\b{re.escape(name)}\b", head):
            return code
    for name, code in ((n.lower(), c) for c, n in MS_NAMES.items()):
        if re.search(rf"\b{re.escape(name)}\b", head):
            return code
    m = re.search(r"(?:^|\n)\s*(" + "|".join(MS_CODES) + r")\s*(?:\n|\s)", "\n".join(pages[:1]))
    return m.group(1).upper() if m else ""


def _decouper_en_blocs(texte: str, cible: int = 900, maxi: int = 1600) -> list[str]:
    """Découpe un texte en blocs d'environ `cible` caractères.

    On coupe d'abord sur les lignes vides quand il y en a, sinon sur les fins
    de phrase. C'est le point qui manquait : le texte extrait d'un PDF ne
    contient pratiquement jamais de ligne vide, si bien qu'un découpage fondé
    sur elles produisait **un segment par page**. Un extrait valait alors une
    page entière, et la citation renvoyée n'était plus localisable.
    """
    paragraphes = [b.strip() for b in re.split(r"\n\s*\n", texte) if b.strip()]
    morceaux: list[str] = []
    for para in paragraphes:
        if len(para) <= maxi:
            morceaux.append(para)
            continue
        # Coupe aux fins de phrase : un extrait cité doit rester retrouvable
        # mot pour mot dans le document source.
        phrases = re.split(r"(?<=[.;:!?])\s+(?=[A-ZÀ-Ý0-9(])", para)
        courant = ""
        for phrase in phrases:
            if courant and len(courant) + len(phrase) + 1 > cible:
                morceaux.append(courant.strip())
                courant = phrase
            else:
                courant = f"{courant} {phrase}".strip()
        if courant.strip():
            morceaux.append(courant.strip())
    return [m for m in morceaux if m]


def split_free_text(pages: list[str], min_chars: int = 320,
                    max_chars: int = 1600) -> list[Segment]:
    """Découpe par paragraphe, ou par phrase quand le document n'en a pas.

    On ne coupe jamais au milieu d'une phrase : un extrait cité doit pouvoir
    être retrouvé tel quel dans le document source.
    """
    author = guess_author(pages)
    segments: list[Segment] = []
    order = 0

    for page_no, texte in enumerate(pages, start=1):
        if not texte.strip():
            continue
        blocs = _decouper_en_blocs(texte, cible=900, maxi=max_chars)
        # Regroupe les blocs trop courts pour porter une citation utile
        fusionnes: list[str] = []
        for bloc in blocs:
            if fusionnes and len(fusionnes[-1]) < min_chars:
                fusionnes[-1] = f"{fusionnes[-1]}\n{bloc}"
            else:
                fusionnes.append(bloc)
        for bloc in fusionnes:
            order += 1
            segments.append(Segment(
                section_id=f"p_{order:04d}",
                section_label=f"Page {page_no}, extrait {order}",
                section_kind="paragraphe",
                text=bloc,
                page_start=page_no,
                page_end=page_no,
                ms_code=author,
                order=order,
            ))
    return segments


def _subdiviser_articles(segments: list[Segment], maxi: int = 2500) -> list[Segment]:
    """Redécoupe les sections trop longues, en conservant leur numéro.

    Un article de dix mille caractères remonte à la recherche comme un bloc
    unique et sature le contexte envoyé au modèle. On le scinde en parties
    numérotées — « Article 5 (2/4) » — pour que la citation pointe la bonne
    portion.

    Les considérants et l'exposé des motifs subissent le même sort : sur une
    proposition de la Commission, le préambule dépassait cent mille
    caractères et occupait à lui seul une part écrasante de l'index.
    """
    DECOUPABLES = ("article", "recital", "expose", "annexe")
    sortie: list[Segment] = []
    for seg in segments:
        if seg.section_kind not in DECOUPABLES or len(seg.text) <= maxi:
            sortie.append(seg)
            continue
        blocs = _decouper_en_blocs(seg.text, cible=maxi * 0.7, maxi=maxi)
        total = len(blocs)
        for i, bloc in enumerate(blocs, start=1):
            sortie.append(Segment(
                section_id=f"{seg.section_id}_{i}" if total > 1 else seg.section_id,
                section_label=(f"{seg.section_label} ({i}/{total})"
                               if total > 1 else seg.section_label),
                section_kind=seg.section_kind,
                text=bloc,
                page_start=seg.page_start,
                page_end=seg.page_end,
                ms_code=seg.ms_code,
                order=seg.order,
            ))
    return sortie


# ---------------------------------------------------------------------------
# Point d'entrée
# ---------------------------------------------------------------------------

@dataclass
class IngestResult:
    kind: str
    segments: list[Segment] = field(default_factory=list)
    author: str = ""
    n_pages: int = 0
    warnings: list[str] = field(default_factory=list)


def ingest(path: str | Path, kind: str | None = None) -> IngestResult:
    """Découpe un document. `kind` force le type si la détection se trompe."""
    path = Path(path)
    kind = kind or detect_kind(path)

    if kind == "wk_table":
        contribs = merge_contributions(parse_wk_pdf(path))
        if not contribs:
            # la détection s'est trompée : on rebascule sans échouer
            pages = read_pages(path)
            res = IngestResult("legal_text", split_legal_text(pages),
                               n_pages=len(pages))
            res.warnings.append(
                "Aucune contribution d'État membre détectée : le document a été "
                "traité comme un texte réglementaire. Corrigez le type si besoin."
            )
            return res
        segments = [
            Segment(
                section_id=c.section_id, section_label=c.section_label,
                section_kind=c.section_kind, text=c.text,
                page_start=c.page_start, page_end=c.page_end,
                ms_code=c.ms_code, contribution_kind=c.kind, order=i,
            )
            for i, c in enumerate(contribs)
        ]
        return IngestResult("wk_table", segments,
                            n_pages=max(c.page_end for c in contribs))

    pages = read_pages(path)
    if kind == "legal_text":
        segments = split_legal_text(pages)
        return IngestResult("legal_text", segments, n_pages=len(pages))

    segments = split_free_text(pages)
    return IngestResult("free_text", segments,
                        author=segments[0].ms_code if segments else "",
                        n_pages=len(pages))


# ---------------------------------------------------------------------------
# Repère de version : nature et date
# ---------------------------------------------------------------------------
#
# Comparer deux textes suppose de savoir lequel précède l'autre, et à quel
# stade de la procédure chacun se situe. « Version antérieure / version
# nouvelle » ne dit ni l'un ni l'autre : une proposition de la Commission et un
# compromis de présidence ne se comparent pas comme deux compromis successifs.
#
# La nature et la date sont donc détectées à l'import — et modifiables, parce
# qu'une détection automatique se trompe et qu'un repère faux vaut moins que
# pas de repère.

VERSION_KINDS = {
    "com": "Proposition de la Commission",
    "compromis": "Compromis de la présidence",
    "orientation": "Orientation générale du Conseil",
    "mandat_pe": "Position du Parlement européen",
    "trilogue": "Accord provisoire (trilogue)",
    "jo": "Texte publié au Journal officiel",
    "autre": "Autre ou non précisé",
}

_INDICES_VERSION = [
    ("jo", [r"journal officiel de l'union", r"official journal of the european",
            r"\bOJ L\b", r"\bJO L\b"]),
    ("trilogue", [r"provisional agreement", r"accord provisoire",
                  r"\btrilogue\b", r"four[- ]column"]),
    ("orientation", [r"general approach", r"orientation g[ée]n[ée]rale"]),
    ("mandat_pe", [r"european parliament.{0,40}(?:mandate|position|report)",
                   r"committee on industry", r"\bIMCO\b", r"\bLIBE\b",
                   r"amendments adopted by the european parliament"]),
    ("compromis", [r"presidency compromise", r"compromise (?:text|proposal)",
                   r"texte de compromis", r"revised compromise",
                   r"compromis de la pr[ée]sidence"]),
    ("com", [r"proposal for a\s+(?:regulation|directive)", r"COM\(\d{4}\)\s*\d+",
             r"proposition de r[èe]glement", r"explanatory memorandum"]),
]

_MOIS = {
    "january": 1, "janvier": 1, "february": 2, "février": 2, "fevrier": 2,
    "march": 3, "mars": 3, "april": 4, "avril": 4, "may": 5, "mai": 5,
    "june": 6, "juin": 6, "july": 7, "juillet": 7, "august": 8, "août": 8,
    "aout": 8, "september": 9, "septembre": 9, "october": 10, "octobre": 10,
    "november": 11, "novembre": 11, "december": 12, "décembre": 12,
    "decembre": 12,
}
_DATE_TEXTE = re.compile(
    r"\b(\d{1,2})\s+([A-Za-zéûôà]+)\s+((?:19|20)\d{2})\b")
_DATE_NUM = re.compile(r"\b(\d{1,2})[./](\d{1,2})[./]((?:19|20)\d{2})\b")


def detect_version_kind(pages: list[str]) -> str:
    """Nature du document dans la procédure, devinée depuis son en-tête.

    On ne regarde que les premières pages : les mentions de procédure figurent
    en tête, et un renvoi à la proposition de la Commission au milieu d'un
    compromis ne doit pas requalifier le document.
    """
    tete = "\n".join(pages[:3]).lower()
    for kind, motifs in _INDICES_VERSION:
        if any(re.search(m, tete, re.IGNORECASE) for m in motifs):
            return kind
    return "autre"


def detect_version_date(pages: list[str]) -> str:
    """Date du document au format ISO (AAAA-MM-JJ), ou chaîne vide.

    On ne regarde que le haut de la première page : sur un document du Conseil
    ou de la Commission, la date d'établissement y figure. Les dates du corps
    du texte — délais, entrée en application, échéances de transposition —
    diraient tout autre chose, et elles sont nombreuses.
    """
    tete = (pages[0] if pages else "")[:1500]

    for m in _DATE_TEXTE.finditer(tete):
        mois = _MOIS.get(m.group(2).lower())
        if not mois:
            continue
        jour, annee = int(m.group(1)), int(m.group(3))
        if 1 <= jour <= 31:
            return f"{annee:04d}-{mois:02d}-{jour:02d}"

    m = _DATE_NUM.search(tete)
    if m:
        jour, mois, annee = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= jour <= 31 and 1 <= mois <= 12:
            return f"{annee:04d}-{mois:02d}-{jour:02d}"
    return ""
